"""
Servicio de procesamiento de documentos.
Extrae texto de archivos PDF y DOCX con metadata.
"""

import os
import logging
from typing import Optional
from dataclasses import dataclass, field

from pypdf import PdfReader
from docx import Document

logger = logging.getLogger(__name__)


@dataclass
class DocumentResult:
    """Resultado de la extracción de un documento."""
    texto: str
    nombre: str
    tipo: str
    paginas: int
    caracteres: int
    metadata: dict = field(default_factory=dict)


class DocumentProcessor:
    """Extrae texto de documentos PDF y DOCX."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

    def process(self, file_path: str, original_filename: Optional[str] = None) -> DocumentResult:
        """
        Procesa un documento y extrae su texto.
        
        Args:
            file_path: Ruta al archivo temporal o permanente
            original_filename: Nombre original del archivo (si es upload)
            
        Returns:
            DocumentResult con texto extraído y metadata
        """
        filename = original_filename or os.path.basename(file_path)
        extension = os.path.splitext(filename)[1].lower()

        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Formato no soportado: {extension}. "
                f"Formatos válidos: {', '.join(self.SUPPORTED_EXTENSIONS)}"
            )

        logger.info(f"Procesando documento: {filename} ({extension})")

        if extension == ".pdf":
            return self._process_pdf(file_path, filename)
        elif extension == ".docx":
            return self._process_docx(file_path, filename)

    def _process_pdf(self, file_path: str, filename: str) -> DocumentResult:
        """Extrae texto de un archivo PDF página por página."""
        reader = PdfReader(file_path)
        pages_text = []

        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                pages_text.append({
                    "pagina": i + 1,
                    "texto": text.strip()
                })

        full_text = "\n\n".join(p["texto"] for p in pages_text)

        return DocumentResult(
            texto=full_text,
            nombre=filename,
            tipo="pdf",
            paginas=len(reader.pages),
            caracteres=len(full_text),
            metadata={
                "paginas_con_texto": len(pages_text),
                "paginas_totales": len(reader.pages),
                "paginas_detalle": pages_text
            }
        )

    def _process_docx(self, file_path: str, filename: str) -> DocumentResult:
        """Extrae texto de un archivo DOCX párrafo por párrafo."""
        doc = Document(file_path)
        paragraphs = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # También extraer texto de tablas
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)

        return DocumentResult(
            texto=full_text,
            nombre=filename,
            tipo="docx",
            paginas=len(paragraphs),  # Aproximación: párrafos como "páginas"
            caracteres=len(full_text),
            metadata={
                "parrafos": len(paragraphs),
                "tablas": len(doc.tables)
            }
        )
